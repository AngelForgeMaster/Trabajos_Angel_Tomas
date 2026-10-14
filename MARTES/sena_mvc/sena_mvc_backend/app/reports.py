from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import (
    UsuarioModel,
    DepartamentoModel,
    DireccionModel,
    ProyectoModel,
)

ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")
ENCABEZADO_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")


def _construir_libro(titulo: str, encabezados: list[str], filas: list[list]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo

    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = ENCABEZADO_FONT
        celda.fill = ENCABEZADO_FILL

    for fila in filas:
        ws.append(fila)

    for indice, encabezado in enumerate(encabezados, start=1):
        columna = get_column_letter(indice)
        ancho = max(len(str(encabezado)), 12)
        for celda in ws[columna]:
            if celda.value is not None:
                ancho = max(ancho, len(str(celda.value)))
        ws.column_dimensions[columna].width = ancho + 2

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_excel_usuarios(usuarios: list[UsuarioModel]) -> BytesIO:
    encabezados = ["ID", "Nombre", "Email", "Rol", "Departamento", "Direccion"]
    filas = [
        [
            usuario.id,
            usuario.nombre,
            usuario.email,
            usuario.rol,
            usuario.departamento.nombre if usuario.departamento else "",
            usuario.departamento.direccion.nombre
            if usuario.departamento and usuario.departamento.direccion
            else "",
        ]
        for usuario in usuarios
    ]
    return _construir_libro("Usuarios", encabezados, filas)


def generar_excel_departamentos(departamentos: list[DepartamentoModel]) -> BytesIO:
    encabezados = ["ID", "Nombre", "Codigo", "Direccion"]
    filas = [
        [
            depto.id,
            depto.nombre,
            depto.codigo,
            depto.direccion.nombre if depto.direccion else "",
        ]
        for depto in departamentos
    ]
    return _construir_libro("Departamentos", encabezados, filas)


def generar_excel_direcciones(direcciones: list[DireccionModel]) -> BytesIO:
    encabezados = ["ID", "Nombre", "Cantidad de Departamentos"]
    filas = [
        [direccion.id, direccion.nombre, len(direccion.departamentos)]
        for direccion in direcciones
    ]
    return _construir_libro("Direcciones", encabezados, filas)


def generar_excel_proyectos(proyectos: list[ProyectoModel]) -> BytesIO:
    encabezados = ["ID", "Nombre", "Descripcion", "Cantidad de Empleados"]
    filas = [
        [proyecto.id, proyecto.nombre, proyecto.descripcion or "", len(proyecto.empleados)]
        for proyecto in proyectos
    ]
    return _construir_libro("Proyectos", encabezados, filas)
